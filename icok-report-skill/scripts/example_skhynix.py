# -*- coding: utf-8 -*-
"""
워크드 예제 — SK하이닉스 밸류에이션 자료 2슬라이드 생성(실제 산출물 재현).
이 스크립트는 icok_style + pptx_build 를 어떻게 엮는지 보여주는 "복붙 출발점"이다.
다른 기업에 적용할 때:
  1) 아래 정합표(DATA)만 교체 → 차트 자동 갱신
  2) build_slide() 안의 "템플릿 구조 의존" 부분(자료 프레임/사이드바 라벨 식별 문자열)을
     팀원 deck의 슬라이드 구조에 맞게 조정
  3) 삽입 위치(move_slide 의 dst index) 조정

실행:
    set PYTHONUTF8=1
    python example_skhynix.py
"""
import os
import sys
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Inches, Emu
import icok_style as S
import pptx_build as B

S.setup()
OUT = os.path.dirname(__file__)

# ── 1) 정합표(DATA) — 모든 자료는 여기서만 인용 (DESIGN_SPEC §10.2) ────────
# 월별 수정주가 / BPS (FnGuide DataGuide6, A000660)
MONTH = [("2021-12",131000,90394),("2022-06",91000,91199),("2022-12",75000,92004),
         ("2023-06",115200,84878),("2023-12",141500,77752),("2024-06",236500,92504),
         ("2024-12",173900,107256),("2025-06",292000,140897),("2025-12",651000,174539),
         ("2026-06",2628000,328204)]   # ※ 실제 산출은 전월 데이터 사용. 여기선 요약본.
dts = [m[0] for m in MONTH]
pbr = [m[1] / m[2] for m in MONTH]

YRS  = ["'21", "'22", "'23", "'24", "'25", "26F", "27F"]
EPS  = [13190, 3063, -12517, 27182, 58955, 325000, 441000]
PE   = [9.9, 24.5, None, 6.4, 11.0, 7.3, 5.4]

# ── 2) 차트 생성 (icok_style) ─────────────────────────────────────────────
c30 = S.band_chart(os.path.join(OUT, "ex_자료30_PBR.png"), dts, pbr, 0.8, 3.8,
                   median=2.3, current="현재 8.0x", ylabel="PBR (배)", ymax=9,
                   note="선행 BPS 4.8x")
c31 = S.bar_line(os.path.join(OUT, "ex_자료31_PER.png"), YRS, EPS, PE,
                 bar_label="EPS (만원)", line_label="Fwd P/E (배)", bar_unit=10000,
                 target=8.5, target_label="목표 8.5x", neg_note="'23 적자 NM",
                 neg_note_xy=(2, -5.6), bar_ylim=(-10, 54), line_ylim=(0, 28),
                 line_offsets={0:(0,-14),1:(0,10),3:(0,-14),4:(0,11),5:(0,-14),6:(0,-14)},
                 target_xy=(2.6, 13.5))
c32 = S.factor_table(os.path.join(OUT, "ex_자료32_discount.png"),
                     [("① 메모리 사이클", "잔존*", "keep"),
                      ("② 코리아 디스카운트", "잔존", "keep"),
                      ("③ 자본집약도(CapEx)", "잔존", "keep"),
                      ("④ 고객집중(NVIDIA)", "잔존", "keep"),
                      ("⑤ 접근성(ADR 상장)", "해소 ↑", "resolve"),
                      ("⑥ HBM 1위(미스프라이싱)", "해소 ↑", "resolve")],
                     title="디스카운트 6요인",
                     footnote="잔존(정당) 4 · 해소 2   *LTA로 일부 완화")
c33 = S.waterfall(os.path.join(OUT, "ex_자료33_waterfall.png"),
                  ["현 멀티플", "+LTA\n가시성", "+HBM4\n재인식", "+ADR\n접근성", "목표"],
                  start=6.5, steps=[1.0, 0.5, 0.5], target=8.5, ceiling=9.0)
print("charts:", [os.path.basename(p) for p in (c30, c31, c32, c33)])

# ── 3) 슬라이드 조립 (pptx_build) — 템플릿 슬라이드#7을 자료 마스터로 복제 ──
DL = r"C:\Users\jwlee\Downloads"
SRC = os.path.join(DL, "ICOK_보고서_SK하이닉스_000660.pptx")   # 베이스 deck
EXHIBIT_MASTER = 7   # 표준 자료 레이아웃 슬라이드 인덱스(팀원 deck에 맞게 조정)


def build_slide(prs, sec, sidelabels, blocks, titles, sources, imgs):
    s = B.duplicate_slide(prs, EXHIBIT_MASTER)
    for sh in list(s.shapes):
        # 자료 캡션 프레임(아래쪽 표) 식별 → 타이틀·출처 채우고 하단으로
        if sh.has_table and "자료 11" in sh.table.cell(0, 0).text:
            B.set_caption(sh.table, titles, sources)
            sh.top = Inches(B.FRAME_TOP)
        elif sh.has_table and "자료 10" in sh.table.cell(0, 0).text:
            sh._element.getparent().remove(sh._element)   # 위쪽 자료 제거
        elif sh.has_text_frame:
            tx = sh.text_frame.text.strip()
            if tx == "기업분석":
                sh.text_frame.paragraphs[0].runs[0].text = sec
            elif "기업 개요" in tx:
                B.fill_body(sh.text_frame, blocks)
            elif "메모리 전환" in tx.replace('\n', ' '):
                B.set_tf(sh.text_frame, sidelabels[0])
            elif "핵심 동력" in tx.replace('\n', ' '):
                B.set_tf(sh.text_frame, sidelabels[1])
            elif tx.startswith("단위"):
                sh.text_frame.paragraphs[0].runs[0].text = ""
            elif "21→30%" in tx:
                sh._element.getparent().remove(sh._element)
    # 위쪽 자료 잔여 도형(화살표 등) 정리
    for sh in list(s.shapes):
        tp = Emu(sh.top).inches if sh.top is not None else -99
        if 5.3 < tp < 7.55 and not sh.has_table and (
                not sh.has_text_frame or not sh.text_frame.text.strip()):
            sh._element.getparent().remove(sh._element)
    B.place_exhibits(s, imgs)
    return s


if not os.path.exists(SRC):
    print(f"[skip 조립] 베이스 deck 없음: {SRC}\n차트만 생성됨. deck 경로를 맞춘 뒤 재실행.")
    sys.exit(0)

p = Presentation(SRC)
build_slide(p, "밸류에이션", ["PBR→PER", "피크이익 역설"],
            [("4.2 밸류에이션 논거 (1) — 왜 PER인가", 'h'),
             ("메모리는 이익 변동성이 커 전통적으로 PBR로 평가돼 왔다. SK하이닉스의 역사적 "
              "PBR 밴드는 0.8~3.8x였으나 현 PBR은 선행 BPS 기준 약 4.8x·당기 8.0x로 역사 "
              "밴드 상단을 명백히 돌파했다(자료 30). LTA가 감익기 OPM을 30%대에서 방어하며 "
              "이익 변동성이 축소되자, 평가 기준이 자산가치(PBR)에서 이익가치(PER)로 이동하는 "
              "것이 정합적이다.", 'b'),
             ("메모리주는 통상 이익 정점에 낮은 P/E가 붙는다(피크이익=저PER). Forward P/E는 "
              "이익저점 2022년 24.5x까지 확대됐다가 2023년 적자로 NM을 거쳐, 이익이 급증하는 "
              "26F·27F에는 7.3x·5.4x로 압축됐다(자료 31). 목표 8.5x는 US 테크식(20x+)이 아니라 "
              "사이클 정상 레인지 상단으로, 피크이익 과대평가 위험을 멀티플에서 차단한 보수적 "
              "값이다.", 'b')],
            ["자료 30. PBR 밴드 (역사 0.8~3.8x 돌파)", "자료 31. 이익급증 ↔ Forward P/E 압축"],
            ["출처: FnGuide DataGuide6, ICOK", "출처: FnGuide DataGuide6, ICOK"],
            [c30, c31])
build_slide(p, "밸류에이션", ["할인 6요인", "Re-rating"],
            [("4.2 밸류에이션 논거 (2) — 할인의 차등 해소", 'h'),
             ("SK하이닉스는 HBM 글로벌 1위·OPM 72%로 Micron(12MF P/E 8~10x)을 앞서는데도 "
              "6~7x로 더 싸게 거래된다. 이 할인을 6개 요인으로 분해하면(자료 32) 정당해 잔존하는 "
              "부분과 좁혀지는 부분이 구분된다. ①메모리 사이클·②코리아 디스카운트·③자본집약도·"
              "④고객집중은 잔존하고, ⑤접근성(ADR)·⑥HBM 1위 미스프라이싱만 좁혀진다.", 'b'),
             ("따라서 8.5x는 '할인 소멸'이 아니라 '좁혀지는 할인분만 반영한 보수적 수렴'이다. "
              "현 약 6.5x(12MF)에서 LTA 이익가시성(+1.0x)·HBM4 재인식(+0.5x)·ADR 접근성(+0.5x)을 "
              "더해 8.5x에 도달하며, 정당한 잔존 할인은 그대로 둔다(자료 33). 천장은 9.0x.", 'b')],
            ["자료 32. 디스카운트 6요인 분해", "자료 33. Re-rating 워터폴 (6.5→8.5x)"],
            ["출처: ICOK", "출처: ICOK"],
            [c32, c33])
# 두 새 슬라이드를 4.3 앞(인덱스 16,17)으로 이동
B.move_slide(p, -2, 16)
B.move_slide(p, -1, 17)
DST = os.path.join(OUT, "ex_SKHynix_valuation_slides.pptx")
p.save(DST)
print("saved:", DST, "| fonts:", B.verify_fonts(DST))
