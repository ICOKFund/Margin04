# -*- coding: utf-8 -*-
"""
ICOK 리포트 PPTX 조립 모듈 (DESIGN_SPEC §5·6·7)
- 핵심: 백지에서 새로 만들지 말고 "템플릿 슬라이드를 복제"한다 → 폰트(Pretendard
  임베딩)·러닝헤더·여백 일관성 유지. 폰트는 프레젠테이션 레벨이라 복제 시 자동 보존.
- 네이티브 차트(has_chart)는 복제 시 rels 직렬화 오류 → 건너뛰고 이미지(PNG)로 대체.

제공 함수:
    duplicate_slide(prs, index)              템플릿 슬라이드 복제
    set_tf(tf, txt) / setcell(cell, txt)     런 단위 텍스트 교체(서식 보존)
    fill_body(tf, blocks)                    본문(소제목=ICOK블루Bold13 / 본문=블랙10.5)
    set_caption(table, titles, sources)      자료 캡션 프레임(3행 표) 타이틀·출처 채우기
    place_exhibits(slide, imgs, ...)         자료 이미지 하단 2-up 배치
    move_slide(prs, src_index, dst_index)    섹션 중간 위치로 슬라이드 이동
    verify_fonts(path)                       임베딩 폰트(.fntdata) 검증
"""
import copy
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from PIL import Image

# 색상 토큰(PPTX) — icok_style 와 동일 (DESIGN_SPEC §2)
ICOK  = RGBColor(0x00, 0x5E, 0xB8)   # 소제목·캡션 헤더·BUY
BLACK = RGBColor(0x1A, 0x1A, 0x1A)   # 본문
GREY  = RGBColor(0x6B, 0x74, 0x80)   # 출처·단위

REFS = (qn('r:embed'), qn('r:id'), qn('r:link'))

# 자료 하단 배치 기본 좌표 (DESIGN_SPEC §6, 단위 in)
FRAME_TOP = 7.68     # 캡션 프레임 top
IMG_TOP   = 7.99     # 이미지 top
IMG_LEFTS = (0.42, 3.80)
IMG_W     = 3.29


def duplicate_slide(prs, index):
    """index 슬라이드를 복제해 맨 끝에 추가하고 반환.
    이미지 rel만 새 rId로 재연결하고 XML 내 참조를 remap. 네이티브 차트는 skip.
    주의: _Relationships._add_relationship 의 3번째 인자는 is_external 라 rId 직접
    지정 불가 → relate_to(새 rId 자동발급) + XML remap 방식 사용."""
    src = prs.slides[index]
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    idmap = {}
    for rId, rel in src.part.rels.items():
        if rel.is_external or not rel.reltype.endswith("/image"):
            continue
        idmap[rId] = new.part.relate_to(rel.target_part, rel.reltype)
    for sh in src.shapes:
        if sh.has_chart:                       # 네이티브 차트 → 이미지로 대체
            continue
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for a in REFS:
                v = node.get(a)
                if v in idmap:
                    node.set(a, idmap[v])
        new.shapes._spTree.append(el)
    return new


def set_tf(tf, txt):
    """text_frame 첫 단락 첫 런의 텍스트만 교체(나머지 런/단락 제거, 서식 유지)."""
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = txt
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run().text = txt
    for ex in tf.paragraphs[1:]:
        ex._p.getparent().remove(ex._p)


def setcell(cell, txt):
    set_tf(cell.text_frame, txt)


def fill_body(tf, blocks):
    """본문 채우기. blocks = list[(text, kind)] / kind: 'h'=소제목 'b'=본문.
    소제목 = Pretendard ICOK블루 Bold 13pt / 본문 = 블랙 10.5pt 줄간격 1.22."""
    tf.word_wrap = True
    tf.clear()
    first = True
    for text, kind in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = "Pretendard"
        if kind == 'h':
            f.size = Pt(13); f.bold = True; f.color.rgb = ICOK
            p.space_before = Pt(2); p.space_after = Pt(5)
        else:
            f.size = Pt(10.5); f.bold = False; f.color.rgb = BLACK
            p.space_after = Pt(7); p.line_spacing = 1.22


def set_caption(table, titles, sources, frame_top=FRAME_TOP):
    """자료 캡션 프레임(3행 표): row0=타이틀(좌/우), row2=출처(좌/우).
    titles/sources = (left, right). 출처 형식: '출처: FnGuide DataGuide6, ICOK' 등."""
    setcell(table.cell(0, 0), titles[0])
    setcell(table.cell(0, 1), titles[1])
    setcell(table.cell(2, 0), sources[0])
    setcell(table.cell(2, 1), sources[1])


def place_exhibits(slide, imgs, lefts=IMG_LEFTS, top=IMG_TOP, width=IMG_W):
    """자료 이미지(PNG)를 페이지 하단에 2-up 배치. 높이는 종횡비로 자동 계산."""
    for img, L in zip(imgs, lefts):
        iw, ih = Image.open(img).size
        slide.shapes.add_picture(img, Inches(L), Inches(top),
                                 width=Inches(width), height=Inches(width * ih / iw))


def move_slide(prs, src_index, dst_index):
    """맨 끝에 추가된 슬라이드를 섹션 중간(dst_index)으로 이동."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    node = ids[src_index]
    lst.remove(node)
    lst.insert(dst_index, node)


def verify_fonts(path):
    """임베딩 폰트(.fntdata) 크기 목록 반환. 완전 임베드면 MB 단위 값이 보인다."""
    import zipfile
    z = zipfile.ZipFile(path)
    sizes = [z.getinfo(n).file_size for n in z.namelist() if n.endswith(".fntdata")]
    return sizes
