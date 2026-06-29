# -*- coding: utf-8 -*-
"""삼성SDI deck의 네이티브 Office 차트 6개를 ICOK 양식 PNG로 교체.
DESIGN_SPEC §7: 네이티브 차트는 이미지로 대체. 같은 앵커(left/top/width)에 배치하고
높이는 종횡비로 자동 계산해 캡션 바와 정렬을 유지한다.
"""
from pptx import Presentation
from pptx.util import Inches, Emu
from PIL import Image

SRC = "/home/user/Margin04/ICOK_삼성SDI_006400.pptx"
AST = "/home/user/Margin04/sdi_assets/"

# (slide_index, left_position_approx_in) -> image
MAP = {
    (9, 0.5): "ja13.png", (9, 3.9): "ja14.png",
    (10, 3.9): "ja16.png",
    (11, 3.9): "ja18.png",
    (12, 0.5): "ja19.png", (12, 3.9): "ja20.png",
}


def pick(slide_idx, left_in):
    best, bestd = None, 99
    for (si, lx), img in MAP.items():
        if si != slide_idx:
            continue
        d = abs(lx - left_in)
        if d < bestd:
            best, bestd = img, d
    return best


prs = Presentation(SRC)
swapped = 0
for i, s in enumerate(prs.slides):
    charts = [sh for sh in s.shapes if sh.has_chart]
    for sh in charts:
        L = Emu(sh.left).inches
        T = Emu(sh.top).inches
        W = Emu(sh.width).inches
        img = pick(i, L)
        if not img:
            print(f"  [WARN] no image for slide {i} chart at L={L:.2f}")
            continue
        iw, ih = Image.open(AST + img).size
        h = W * ih / iw
        # 네이티브 차트 graphicFrame 제거
        sh._element.getparent().remove(sh._element)
        # 같은 앵커에 PNG 배치 (top-align, 폭 동일, 높이는 종횡비)
        s.shapes.add_picture(AST + img, Inches(L), Inches(T), Inches(W), Inches(h))
        print(f"  slide {i}: {img}  @({L:.2f},{T:.2f}) W={W:.2f} H={h:.2f}")
        swapped += 1

print("swapped:", swapped)
prs.save(SRC)
print("saved:", SRC)

# 검증: 차트가 모두 사라지고 그림으로 바뀌었는지
chk = Presentation(SRC)
nchart = sum(1 for s in chk.slides for sh in s.shapes if sh.has_chart)
npic = sum(1 for s in chk.slides for sh in s.shapes if sh.shape_type == 13)
print(f"verify: remaining native charts={nchart}, pictures={npic}")
