#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Micron F3Q26 Preview - 증권사 비교 차트 Excel 생성 (ICOK 네이비 팔레트)."""
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import Marker
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice

# ICOK 네이비 그라데이션 8색 (매크로와 동일)
PALETTE = ["002A52", "003F7B", "00509D", "2E76C1",
           "5C8DCA", "7399CE", "8AA5D3", "ABB9DB"]
NAVY = "002A52"
GREY = "BFBFBF"

HDR_FILL = PatternFill("solid", fgColor="002A52")
HDR_FONT = Font(name="Nanum Gothic", color="FFFFFF", bold=True, size=10)
CELL_FONT = Font(name="Nanum Gothic", size=10)
BOLD = Font(name="Nanum Gothic", size=10, bold=True)
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CTR = Alignment(horizontal="center", vertical="center")


def style_table(ws, top, left, nrows, ncols):
    for r in range(top, top + nrows):
        for c in range(left, left + ncols):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = CTR
            if r == top:
                cell.fill = HDR_FILL
                cell.font = HDR_FONT
            else:
                cell.font = CELL_FONT


def set_series_colors(chart, line=False):
    for i, s in enumerate(chart.series):
        col = PALETTE[i % len(PALETTE)]
        if line:
            s.graphicalProperties = GraphicalProperties()
            s.graphicalProperties.line = LineProperties(solidFill=col, w=19050)  # 1.5pt
            s.smooth = False
        else:
            s.graphicalProperties = GraphicalProperties(solidFill=col)
            s.graphicalProperties.line = LineProperties(noFill=True)


wb = Workbook()

# ============================================================
# 자료1. 연간 실적 추이 (매출 bar + GM% line)
# ============================================================
ws = wb.active
ws.title = "자료1_연간실적"
data = [
    ["항목(FY)", "2024A", "2025A", "2026E", "2027E", "2028E"],
    ["매출액($mn)", 25111, 37378, 115003, 197500, 210000],
    ["GM%", 22.4, 39.8, 76.9, 82.9, 80.7],
]
for r, row in enumerate(data, 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
style_table(ws, 1, 1, 3, 6)

bar = BarChart()
bar.type = "col"; bar.gapWidth = 25
bar.title = "자료1. Micron 연간 실적 추이 (매출 vs GM%)"
bar.add_data(Reference(ws, min_col=1, min_row=2, max_col=6, max_row=2), titles_from_data=True, from_rows=True)
bar.set_categories(Reference(ws, min_col=2, min_row=1, max_col=6, max_row=1))
bar.y_axis.title = "매출($mn)"
set_series_colors(bar)
bar.y_axis.majorGridlines = None

line = LineChart()
line.add_data(Reference(ws, min_col=1, min_row=3, max_col=6, max_row=3), titles_from_data=True, from_rows=True)
line.y_axis.axId = 200
line.y_axis.title = "GM%"
line.y_axis.crosses = "max"
for s in line.series:
    s.graphicalProperties = GraphicalProperties()
    s.graphicalProperties.line = LineProperties(solidFill="2E76C1", w=22225)
    s.marker = Marker(symbol="circle", size=7)
    s.smooth = False
bar += line
bar.width, bar.height = 20, 10
ws.add_chart(bar, "A6")

# ============================================================
# 자료7. 분기 실적 추이 (매출/OP bar + OPM% line)
# ============================================================
ws = wb.create_sheet("자료7_분기실적")
data = [
    ["분기", "1QFY26", "2QFY26", "3QFY26e", "4QFY26e"],
    ["매출($bn)", 13.6, 23.9, 35.3, 41.7],
    ["영업이익($bn)", 6.4, 16.5, 27.3, 32.4],
    ["OPM%", 47, 69, 77, 78],
]
for r, row in enumerate(data, 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
style_table(ws, 1, 1, 4, 5)

bar = BarChart()
bar.type = "col"; bar.gapWidth = 25
bar.title = "자료7. Micron 분기 실적 추이"
bar.add_data(Reference(ws, min_col=1, min_row=2, max_col=5, max_row=3), titles_from_data=True, from_rows=True)
bar.set_categories(Reference(ws, min_col=2, min_row=1, max_col=5, max_row=1))
bar.y_axis.title = "$bn"
set_series_colors(bar)
bar.y_axis.majorGridlines = None

line = LineChart()
line.add_data(Reference(ws, min_col=1, min_row=4, max_col=5, max_row=4), titles_from_data=True, from_rows=True)
line.y_axis.axId = 200
line.y_axis.title = "OPM%"
line.y_axis.crosses = "max"
for s in line.series:
    s.graphicalProperties = GraphicalProperties()
    s.graphicalProperties.line = LineProperties(solidFill="5C8DCA", w=22225)
    s.marker = Marker(symbol="circle", size=7)
    s.smooth = False
bar += line
bar.width, bar.height = 20, 10
ws.add_chart(bar, "A7")

# ============================================================
# 자료8. F3Q26 GS vs Street (clustered bar)
# ============================================================
ws = wb.create_sheet("자료8_F3Q26비트")
data = [
    ["항목", "GS 추정", "스트리트"],
    ["매출($bn)", 37.6, 34.4],
    ["EPS($)", 22.07, 19.74],
]
for r, row in enumerate(data, 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
style_table(ws, 1, 1, 3, 3)

bar = BarChart(); bar.type = "col"; bar.gapWidth = 25
bar.title = "자료8. F3Q26 GS 추정 vs 스트리트 (Beat)"
bar.add_data(Reference(ws, min_col=2, min_row=1, max_col=3, max_row=3), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=2, max_col=1, max_row=3))
set_series_colors(bar)
bar.y_axis.majorGridlines = None
bar.width, bar.height = 16, 9
ws.add_chart(bar, "A6")

# ============================================================
# 자료9. F4Q26 가이던스 vs Street
# ============================================================
ws = wb.create_sheet("자료9_가이던스")
data = [
    ["항목", "GS 예상", "스트리트"],
    ["매출($bn)", 48.8, 40.4],
    ["EPS($)", 29.95, 23.68],
]
for r, row in enumerate(data, 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
style_table(ws, 1, 1, 3, 3)

bar = BarChart(); bar.type = "col"; bar.gapWidth = 25
bar.title = "자료9. F4Q26 가이던스 전망 vs 스트리트 (Guide-up)"
bar.add_data(Reference(ws, min_col=2, min_row=1, max_col=3, max_row=3), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=2, max_col=1, max_row=3))
set_series_colors(bar)
bar.y_axis.majorGridlines = None
bar.width, bar.height = 16, 9
ws.add_chart(bar, "A6")

# ============================================================
# 자료11. 증권사별 목표주가
# ============================================================
ws = wb.create_sheet("자료11_목표주가")
data = [
    ["증권사", "목표주가($)"],
    ["Citi", 1200],
    ["HSBC", 1100],
    ["Goldman Sachs", 900],
]
for r, row in enumerate(data, 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
style_table(ws, 1, 1, 4, 2)

bar = BarChart(); bar.type = "bar"; bar.gapWidth = 25
bar.title = "자료11. 증권사별 목표주가"
bar.add_data(Reference(ws, min_col=2, min_row=1, max_col=2, max_row=4), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=2, max_col=1, max_row=4))
# 단일 시리즈 → 막대마다 다른 색
from openpyxl.chart.series import DataPoint
s = bar.series[0]
s.graphicalProperties = GraphicalProperties(solidFill=NAVY)
for i in range(3):
    dp = DataPoint(idx=i)
    dp.graphicalProperties = GraphicalProperties(solidFill=PALETTE[i])
    s.data_points.append(dp)
bar.legend = None
bar.width, bar.height = 16, 8
ws.add_chart(bar, "A6")

# ============================================================
# 자료12. 증권사별 EPS 추정 비교
# ============================================================
ws = wb.create_sheet("자료12_EPS비교")
data = [
    ["EPS($)", "FY25A", "FY26E", "FY27E", "FY28E"],
    ["Goldman Sachs", 7.42, 67.48, 138.86, 137.51],
    ["HSBC", 7.59, 61.88, 126.82, 142.91],
    ["Citi", 7.43, 60.73, 114.73, 117.83],
    ["컨센서스", 8.0, 58.0, 106.0, 105.0],
]
for r, row in enumerate(data, 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
style_table(ws, 1, 1, 5, 5)

bar = BarChart(); bar.type = "col"; bar.gapWidth = 25
bar.title = "자료12. 증권사별 EPS 추정 비교 (Non-GAAP)"
bar.add_data(Reference(ws, min_col=1, min_row=2, max_col=5, max_row=5), titles_from_data=True, from_rows=True)
bar.set_categories(Reference(ws, min_col=2, min_row=1, max_col=5, max_row=1))
set_series_colors(bar)
bar.y_axis.majorGridlines = None
bar.y_axis.title = "EPS($)"
bar.width, bar.height = 20, 10
ws.add_chart(bar, "A7")

# 열 너비 정리
for sh in wb.worksheets:
    sh.column_dimensions["A"].width = 18
    for col in "BCDEF":
        sh.column_dimensions[col].width = 12

out = "/home/user/Margin04/Micron_3Q26_Charts.xlsx"
wb.save(out)
print("saved:", out)
