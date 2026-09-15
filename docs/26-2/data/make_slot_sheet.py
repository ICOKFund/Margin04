# -*- coding: utf-8 -*-
"""팀별 담당 산업 배정 시트 (A4 1장 × 3팀) 생성.

사용: python3 make_slot_sheet.py [CSV] [NAV] [출력경로]
"""
import csv, io, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from sector_map import build, MEGA, MANDATE, split_sectors
from weights import load

CSV = sys.argv[1] if len(sys.argv) > 1 else 'krx300_constituents_20260915.csv'
NAV = float(sys.argv[2]) if len(sys.argv) > 2 else 28_000_000
OUT = sys.argv[3] if len(sys.argv) > 3 else 'ICOK_담당산업_배정시트.docx'
ASOF = '2026-09-15'

TEAM = {
    1: ('테크·전동화', '배준서', 6, '반도체가 잘 팔릴 때 · 환율이 오를 때'),
    2: ('중후장대·인프라', '이재우', 5, '공장·발전소 짓는 돈이 늘어날 때'),
    3: ('내수·디펜시브', '조찬형', 5, '금리가 내리고 국내 경기가 살아날 때'),
}
NOTE = {
    1: '담당 산업 6개, 인원 6명으로 1인 1산업이다. 순수지주는 「지주」로 따로 묶었다 — '
       'SK스퀘어는 SK하이닉스 룩스루 배수가 1.8배이므로 편입 시 펀드 전체 노출을 함께 계산한다.',
    2: '담당 산업 6개, 인원 5명이므로 팀장이 1개를 겸임한다. 지주 분리로 작아진 '
       '소재·화학·철강과 에너지·상사·운송을 「소재·에너지」로 합쳤다.',
    3: '담당 산업 6개, 인원 5명이므로 팀장이 1개를 겸임한다. 금융이 팀 몫의 절반이라 '
       '은행·지주 / 보험·증권으로 나눴다. 보험(IFRS17 CSM)과 매출 미발생 바이오는 팀장이 본다.',
}
FONT = '맑은 고딕'


def style(run, size=9, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=9, bold=False, color=None, after=2, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    style(p.add_run(text), size, bold, color)
    return p


def cell(c, text, size=8, bold=False, color=None, align=None):
    c.text = ''
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    style(p.add_run(text), size, bold, color)


def main():
    rows = load(CSV)
    total = sum(float(r[5]) for r in rows)
    smap = build()
    flat = []
    for code, name, close, _, _, mcap in rows:
        if name in MEGA:
            continue
        flat.append(dict(name=name, px=float(close), sec=smap[name],
                         w=float(mcap) / total * 100))
    md = split_sectors(flat)
    S = {}
    for x in flat:
        S.setdefault(x['sec'], []).append((x['name'], x['px'], x['w']))

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for m in ('top_margin', 'bottom_margin'):
        setattr(s, m, Cm(1.3))
    for m in ('left_margin', 'right_margin'):
        setattr(s, m, Cm(1.5))

    for t in (1, 2, 3):
        if t > 1:
            doc.add_section(WD_SECTION.NEW_PAGE)
            s2 = doc.sections[-1]
            s2.page_width, s2.page_height = Cm(21), Cm(29.7)
            s2.top_margin = s2.bottom_margin = Cm(1.3)
            s2.left_margin = s2.right_margin = Cm(1.5)

        name, lead, n, when = TEAM[t]
        secs = sorted([k for k, v in md.items() if v == t],
                      key=lambda k: -sum(x[2] for x in S[k]))
        bw = sum(sum(x[2] for x in S[k]) for k in secs)

        para(doc, f'팀 {t} · {name}', 17, True, after=1)
        para(doc, f'팀장 {lead} · {n}명   |   팀 예산 {bw/100*NAV:,.0f}원 (벤치마크 {bw:.2f}%)   |   '
                  f'{ASOF} 종가 기준 · NAV {NAV:,.0f}원', 9, color=(0x55, 0x55, 0x55), after=3)
        para(doc, f'우리 팀이 잘 되는 때 — {when}', 9, color=(0x55, 0x55, 0x55), after=8)

        para(doc, '담당 산업', 11, True, after=3)
        hdr = ['담당 산업', 'BM비중', '배정액', '종목수', '최저가', '1지망', '2지망']
        tb = doc.add_table(rows=1, cols=len(hdr))
        tb.style = 'Table Grid'
        for i, h in enumerate(hdr):
            cell(tb.rows[0].cells[i], h, 8, True,
                 align=WD_ALIGN_PARAGRAPH.CENTER if i else None)
        for k in secs:
            v = sorted(S[k], key=lambda x: -x[2])
            w = sum(x[2] for x in v)
            r = tb.add_row().cells
            cell(r[0], k, 9, True)
            cell(r[1], f'{w:.2f}%', 9, align=WD_ALIGN_PARAGRAPH.CENTER)
            cell(r[2], f'{w/100*NAV:,.0f}원', 9, align=WD_ALIGN_PARAGRAPH.RIGHT)
            cell(r[3], f'{len(v)}', 9, align=WD_ALIGN_PARAGRAPH.CENTER)
            cell(r[4], f'{min(x[1] for x in v):,.0f}원', 9, align=WD_ALIGN_PARAGRAPH.RIGHT)
            cell(r[5], '', 9); cell(r[6], '', 9)
        for w_, c_ in zip((3.6, 1.6, 2.3, 1.2, 2.0, 1.6, 1.6), range(7)):
            for row in tb.rows:
                row.cells[c_].width = Cm(w_)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        para(doc, '산업별 주요 종목  (벤치마크 비중 상위 6개 · 괄호는 BM비중)', 11, True, after=3)
        tb2 = doc.add_table(rows=0, cols=2)
        tb2.style = 'Table Grid'
        for k in secs:
            v = sorted(S[k], key=lambda x: -x[2])[:6]
            r = tb2.add_row().cells
            cell(r[0], k, 8, True)
            cell(r[1], '  '.join(f'{x[0]} {x[1]:,.0f}원({x[2]:.2f}%)' for x in v), 8)
            r[0].width, r[1].width = Cm(3.6), Cm(14.3)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        para(doc, '이 팀의 배정 원칙', 10, True, after=2)
        para(doc, NOTE[t], 9, after=6)

        para(doc, '과제 — 담당 산업 스크리닝  (A4 1~2장 · 마감 9월 22일 세션 전날 자정)', 10, True, after=2)
        for line in (
            '1.  담당 산업 전 종목을 매출 CAGR(3년) · 영업이익률 · ROE 세 지표로 줄 세운다.',
            '     모든 수치에 출처와 기준일을 기재한다. 1순위 자료는 DART 사업보고서다.',
            '2.  상위 3종목을 선정하고 각각 두 줄로 근거를 기재한다.',
            '3.  탈락 5종목과 사유를 한 줄씩 기재한다. 해당 지표와 수치를 명시한다.',
            '4.  1주 가격이 팀 예산의 20%를 초과하는 종목을 표시한다.',
            '5.  지주를 검토할 때는 같은 팀이 보유한 자회사와의 합산 노출을 함께 확인한다.',
        ):
            para(doc, line, 9, after=1)
        para(doc, '증권사 리포트를 논지의 근거로 사용하지 않는다. 목표주가를 인용하지 않는다.',
             9, True, after=2)
        para(doc, '스크리닝 결과로 팀 미니 IC에서 1라운드 발표 섹터를 정한다.  '
                  '효성중공업(2,689,000원 · 지수 비중 0.47%)은 1주 매수만으로 한도를 초과한다.',
             9, color=(0x99, 0x00, 0x11), after=0)

    doc.save(OUT)
    print(f'저장: {OUT}')


if __name__ == '__main__':
    main()
