# -*- coding: utf-8 -*-
"""9/29 세부섹터 스터디 안내 공문 (A4 세로) + 붙임 1 작성 양식 (A4 가로).

사용: python3 make_notice.py [출력경로]
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from make_sector_study import build, shade

OUT = sys.argv[1] if len(sys.argv) > 1 else 'ICOK_세부섹터_스터디_안내공문.docx'
F = '맑은 고딕'
BLACK = (0, 0, 0); GRAY = (0x55, 0x55, 0x55)
BODY = 10.5
IND = 0.8     # 번호 단계당 들여쓰기(cm)


def run(p, text, size=BODY, bold=False, color=BLACK):
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = F
    r._element.rPr.rFonts.set(qn('w:eastAsia'), F)
    r.font.color.rgb = RGBColor(*color)
    return r


def para(doc, text='', size=BODY, bold=False, level=0, hang=0.0, after=4, before=0,
         align=None, color=BLACK, line=1.35):
    """level: 번호 단계(0=1., 1=가., 2=1)). hang: 내어쓰기 폭(cm)."""
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_after = Pt(after); f.space_before = Pt(before); f.line_spacing = line
    f.left_indent = Cm(IND * level + hang); f.first_line_indent = Cm(-hang)
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, size, bold, color)
    return p


def rule(doc, weight=12, after=4):
    """문단 아래 가로선."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', str(weight)), ('space', '1'), ('color', '000000')):
        bt.set(qn(f'w:{k}'), v)
    b.append(bt); pPr.append(b)
    return p


def cell(c, text, size=9.5, bold=False, align=None, fill=None):
    c.text = ''
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    if align is not None:
        p.alignment = align
    lines = text.split('\n')
    for i, ln in enumerate(lines):
        if i:
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(1)
            if align is not None:
                p.alignment = align
        run(p, ln, size, bold)
    if fill:
        shade(c, fill)


def table(doc, headers, rows, ws, level=1, center=()):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell(t.rows[0].cells[i], h, 9.5, True, WD_ALIGN_PARAGRAPH.CENTER, 'E7E6E6')
    for row in rows:
        r = t.add_row().cells
        for i, v in enumerate(row):
            cell(r[i], v, 9.5, i == 0 and 0 not in center,
                 WD_ALIGN_PARAGRAPH.CENTER if i in center else None)
    for i, w in enumerate(ws):
        for r in t.rows:
            r.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def footer(doc):
    rule(doc, 12, after=3)
    t = doc.add_table(rows=3, cols=6)
    rows = (('기안자', '', '검토자', '', '결재권자', '회장'),
            ('협조자', '팀장 3인', '', '', '', ''),
            ('시행', 'ICOK 26-2-○○호 (2026. 9. 26.)', '', '', '접수', ''))
    for r, vals in zip(t.rows, rows):
        for c, v in zip(r.cells, vals):
            cell(c, v, 8.5, v in ('기안자', '검토자', '결재권자', '협조자', '시행', '접수'))
    for i, w in enumerate((1.8, 4.2, 1.8, 3.4, 1.8, 3.4)):
        for r in t.rows:
            r.cells[i].width = Cm(w)
    rule(doc, 4, after=2)
    para(doc, 'ICOK Fund  ·  문의  회장단  ·  공개구분  대내 공개', 8.5, color=GRAY, after=0)


def notice(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.5)
    s.left_margin = s.right_margin = Cm(2.0)

    # 머리
    para(doc, 'ICOK Fund', 20, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    for k, v in (('수신', '팀 1 · 팀 2 · 팀 3 팀원 전원'),
                 ('참조', '각 팀장'),
                 ('(경유)', '')):
        p = para(doc, after=1)
        run(p, f'{k}  ', BODY, True); run(p, v)
    p = para(doc, after=0, before=2)
    run(p, '제목  ', BODY, True)
    run(p, '9월 29일 정기세션 「세부섹터 스터디 종합본」 작성 및 발표 안내', BODY, True)
    rule(doc, 12, after=8)

    # 1.
    para(doc, '1.  관련', level=0, hang=0.6)
    para(doc, '가.  ICOK Fund 26-2 팀별 담당 산업 배정 (2026. 9. 15.)', level=1, hang=0.6, after=1)
    para(doc, '나.  팀별 담당 산업 스크리닝 취합 결과 (2026. 9. 22.)', level=1, hang=0.6, after=6)

    # 2.
    para(doc, '2.  위 호와 관련하여 팀별로 담당 산업 중 세부섹터 1개를 선정하여 기초 자료를 학습하고, '
              '그 결과를 9월 29일(화) 정기세션에서 발표하고자 하오니 아래 사항을 참고하여 '
              '작성·제출하여 주시기 바랍니다.', level=0, hang=0.6, after=5)

    para(doc, '가.  목적', level=1, hang=0.6, after=1)
    para(doc, '1)  인뎁스 산업 리포트 작성이 아니라, 팀원 전원이 담당 세부섹터의 구조를 공부하는 것을 목적으로 함', level=2, hang=0.6, after=1)
    para(doc, '2)  투자 결론을 요구하지 않으며, 결과물은 10월 6일 종목 발제의 기초 자료로 활용함', level=2, hang=0.6, after=5)

    para(doc, '나.  과제 개요', level=1, hang=0.6, after=3)
    table(doc, ['구분', '내용'], [
        ('대상', '팀 1 · 팀 2 · 팀 3 전원'),
        ('세부섹터', '팀별 1개. 팀 미니 IC에서 자율 선정하고 9. 27.(일)까지 회장단에 공유'),
        ('산출물', '세부섹터 스터디 종합본 1부 (붙임 양식, A4 가로)'),
        ('발표', '팀당 12분 + 질의 5분. 종합본을 화면 공유하여 발표하며 별도 슬라이드는 만들지 않음'),
        ('제출', '2026. 9. 28.(월) 23:59까지 팀장이 회장단에 일괄 제출'),
    ], (3.2, 13.3))

    para(doc, '다.  구성 및 발표 시간', level=1, hang=0.6, after=3)
    table(doc, ['장', '내용', '공부 질문 (요지)', '발표'], [
        ('Ⅰ-1', '산업 정의와 경계', '무엇을 팔고 누가 사는가. 경계는 어디인가', '1분'),
        ('Ⅰ-2', '밸류체인', '몇 단계를 거치는가. 담당 종목은 어느 단계에 있는가', '1.5분'),
        ('Ⅰ-3', '마진 구조', '가격을 쥔 단계와 원가를 떠안는 단계는 어디인가', '1분'),
        ('Ⅰ-4', '경쟁 구도', '마진이 높은 단계에서 누가 이기고 있는가', '1분'),
        ('Ⅰ-5', '지난 사이클', '과거 호황·불황은 무엇이 시작하고 끝냈는가', '0.5분'),
        ('Ⅱ', '수익 구조 · 사이클 변수 3개', '매출은 단가와 물량 중 무엇이 움직이는가', '3.5분'),
        ('Ⅲ', '시장이 기대하는 것', '시장은 어느 종목에 높은 배수를 주는가', '1.5분'),
        ('Ⅳ', '우리가 모르는 것 5개', '답이 나오지 않은 질문과 확인 방법', '2분'),
        ('부록', '용어집 20개', '—', '제외'),
    ], (1.5, 4.8, 8.4, 1.8), center=(0, 3))

    para(doc, '라.  작성 기준', level=1, hang=0.6, after=1)
    for i, ln in enumerate((
            '1순위 자료는 DART 사업보고서 「II. 사업의 내용」으로 하며, 3개사 이상을 비교하여 읽음',
            '모든 수치에 출처와 기준일을 기재함',
            '각 장의 공부 질문에 답하는 것을 완료 기준으로 하며, 서술만 있고 산출이 없는 장은 미제출로 봄',
            '그래프 3개 이상은 팀이 직접 작성함. 캡처 이미지는 출처를 기재하되 3개에 포함하지 않음'), 1):
        para(doc, f'{i})  {ln}', level=2, hang=0.6, after=1)
    para(doc, after=3)

    para(doc, '마.  금지 사항', level=1, hang=0.6, after=1)
    for i, ln in enumerate((
            '증권사 리포트를 논지의 근거로 사용하거나 목표주가를 인용하는 것',
            '검증 불가능한 표현 — 유망하다, 성장성이 크다, 저평가되어 있다',
            '전망 표현 — 전망한다, 수혜가 예상된다, 성장이 기대된다'), 1):
        para(doc, f'{i})  {ln}', level=2, hang=0.6, after=1)
    para(doc, after=3)

    para(doc, '바.  추진 일정', level=1, hang=0.6, after=3)
    table(doc, ['일자', '내용'], [
        ('9. 26.(토)', '세부섹터 선정, 장별 담당 배분, 사업보고서 II장 읽기'),
        ('9. 27.(일)', '장별 작성, 세부섹터 회장단 공유'),
        ('9. 28.(월)', '팀 취합 및 수치 교차검증, Ⅳ장 팀 토의, 23:59 제출'),
        ('9. 29.(화)', '정기세션 발표 (팀당 12분 + 질의 5분)'),
    ], (3.2, 13.3), center=(0,))

    para(doc, '3.  9월 24일 배포한 양식으로 작성 중인 내용은 새 양식의 해당 장으로 옮겨 사용하시기 바랍니다. '
              '기존 「1. 섹터 정의와 밸류체인」은 Ⅰ-1·Ⅰ-2로 나뉘었고, Ⅰ-3 마진 구조와 Ⅰ-5 지난 사이클이 '
              '신설되었으며, 나머지 장은 번호만 바뀌었습니다.', level=0, hang=0.6, after=5)
    para(doc, '4.  장별 담당 배분 권장안과 사업보고서 II장 항목별 대응표는 붙임의 「작업 안내」를 참고하시기 바랍니다.',
         level=0, hang=0.6, after=10)

    p = para(doc, after=16)
    run(p, '붙임  '); run(p, '세부섹터 스터디 종합본 작성 양식 1부.  끝.')

    para(doc, 'ICOK Fund 회장', 16, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    footer(doc)


def main():
    doc = Document()
    notice(doc)
    build(doc, attach=True)
    doc.save(OUT)
    print(f'저장: {OUT}  구역 {len(doc.sections)}개')


if __name__ == '__main__':
    main()
