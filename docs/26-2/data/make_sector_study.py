# -*- coding: utf-8 -*-
"""세부섹터 스터디 종합본 템플릿 (A4 가로) 생성.

9/29 세션용. 화면 공유를 전제로 가로 배치한다.
Ⅰ 산업 구조(5장)를 중심에 두고, 각 장에 공부 질문과 1순위 자료 위치를 붙인다.
사용: python3 make_sector_study.py [출력경로]
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else 'ICOK_세부섹터_스터디_템플릿.docx'
F = '맑은 고딕'
NAVY = (0x1F, 0x2A, 0x44); GRAY = (0x5A, 0x5A, 0x5A); RED = (0x9C, 0x00, 0x06)
W = 26.7   # 본문 폭(cm)


def style(run, size=10, bold=False, color=None):
    run.font.size = Pt(size); run.font.bold = bold; run.font.name = F
    run._element.rPr.rFonts.set(qn('w:eastAsia'), F)
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text='', size=10, bold=False, color=None, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
    style(p.add_run(text), size, bold, color)
    return p


def shade(c, fill):
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), fill)
    c._tc.get_or_add_tcPr().append(el)


def cell(c, text='', size=9, bold=False, color=None, align=None, fill=None):
    c.text = ''
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    style(p.add_run(text), size, bold, color)
    if fill:
        shade(c, fill)


def widths(t, ws):
    for i, w in enumerate(ws):
        for r in t.rows:
            r.cells[i].width = Cm(w)


def table(doc, headers, ws, rows=1, first=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell(t.rows[0].cells[i], h, 9, True, NAVY,
             WD_ALIGN_PARAGRAPH.CENTER if i else None, 'D6DCE4')
    for k in range(rows):
        r = t.add_row().cells
        if first:
            cell(r[0], first[k] if k < len(first) else '', 9, True, NAVY, fill='F2F2F2')
    widths(t, ws)
    return t


def box(doc, text, h):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'
    cell(t.rows[0].cells[0], text, 9, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)
    t.rows[0].cells[0].width = Cm(W); t.rows[0].height = Cm(h)


def page(doc, first=False):
    if not first:
        doc.add_section(WD_SECTION.NEW_PAGE)
    s = doc.sections[-1]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21)
    s.top_margin = s.bottom_margin = Cm(1.1)
    s.left_margin = s.right_margin = Cm(1.5)


def header(doc, no, title, minutes, produce, source, questions):
    para(doc, f'{no}   {title}', 16, True, NAVY, after=1)
    para(doc, f'발표 {minutes}  |  담당자 __________  |  산출: {produce}', 9, color=GRAY, after=1)
    para(doc, f'1순위 자료: {source}', 9, color=GRAY, after=5)
    t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
    cell(t.rows[0].cells[0], '공부 질문', 9, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 'EEF1F6')
    c = t.rows[0].cells[1]; shade(c, 'EEF1F6'); c.text = ''
    for i, q in enumerate(questions):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        style(p.add_run(f'{i+1}.  {q}'), 9)
    widths(t, (2.4, W - 2.4))
    gap(doc, 4)


def gap(doc, pt=5):
    doc.add_paragraph().paragraph_format.space_after = Pt(pt)


def note(doc, text):
    para(doc, text, 9, color=GRAY, after=0)


def build(doc, attach=False):
    """양식 전체를 doc 뒤에 붙인다. attach=True면 공문의 붙임으로 새 구역에서 시작한다."""
    # ── 표지 ────────────────────────────────────────────────────
    page(doc, first=not attach)
    if attach:
        para(doc, '붙임', 10, True, GRAY, after=2)
    para(doc, '세부섹터 스터디 종합본', 20, True, NAVY, after=1)
    para(doc, 'ICOK Fund 26-2  ·  2026년 9월 29일 세션  ·  팀당 발표 12분 + 질의 5분  ·  제출 9월 28일 자정',
         10, color=GRAY, after=8)
    t = table(doc, ['항목', '기재'], (5, W - 5), rows=5,
              first=['팀', '세부섹터', '선정 사유 (한 줄)', '작성자', '작성 기준일'])
    gap(doc, 6)

    para(doc, '발표 진행', 12, True, NAVY, after=3)
    t = table(doc, ['구분', '장', '내용', '발표', '담당자'], (4.5, 2.2, 13, 2.5, 4.5), rows=5,
              first=['Ⅰ  산업 구조', 'Ⅱ  산업을 움직이는 것', 'Ⅲ  시장이 기대하는 것',
                     'Ⅳ  우리가 모르는 것', '부록  용어집'])
    for r, (pg, what, mn) in zip(t.rows[1:], (
            ('5장', '정의와 경계 · 밸류체인 · 마진 구조 · 경쟁 구도 · 지난 사이클', '5분'),
            ('2장', '수익 구조 · 사이클 변수 3개', '3.5분'),
            ('1장', '밸류에이션 산포도 · 컨센서스 변화', '1.5분'),
            ('1장', '미결 질문 5개', '2분'),
            ('1장', '용어 20개', '발표 제외'))):
        cell(r.cells[1], pg, 9, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell(r.cells[2], what, 9)
        cell(r.cells[3], mn, 9, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 6)

    para(doc, '이 과제가 요구하는 것', 12, True, NAVY, after=3)
    for ln in (
        '이 과제는 산업 전망이 아니라 산업 공부다. 투자 결론을 쓰지 않는다.',
        '각 장은 「무엇을 계산하거나 수집해 왔는가」로 평가한다. 서술만 있고 산출이 없으면 미제출로 본다.',
        '각 장의 공부 질문에 답할 수 있으면 그 장은 끝난 것이다. 칸을 채우는 것이 목적이 아니다.',
    ):
        para(doc, f'·  {ln}', 10, after=2)

    # ── 작업 안내 ───────────────────────────────────────────────
    page(doc)
    para(doc, '작업 안내   (발표 전 이 장은 삭제한다)', 16, True, NAVY, after=6)
    para(doc, '첫날 할 일 — 사업보고서 「II. 사업의 내용」을 3개사 이상 읽는다', 12, True, NAVY, after=3)
    note(doc, '사업보고서 II장은 회사가 자기 산업을 직접 설명하는 1차 자료다. 3개사를 나란히 읽으면 서로 다르게 설명하는 지점이 보이고, 그 지점이 공부할 거리다.')
    gap(doc, 3)
    t = table(doc, ['사업보고서 II장 항목', '이 문서에서 쓰는 장'], (10, W - 10), rows=7,
              first=['산업의 특성 · 산업의 성장성', '경기변동의 특성', '시장점유율 · 경쟁요소 · 회사의 경쟁우위요소',
                     '주요 제품 등의 현황 · 매출실적', '주요 원재료 · 매출처(주요 고객)',
                     '관계법령 또는 정부의 규제', '생산능력 · 설비'])
    for r, v in zip(t.rows[1:], ('Ⅰ-1 정의와 경계', 'Ⅰ-5 지난 사이클 · Ⅱ-2 사이클 변수',
                                  'Ⅰ-4 경쟁 구도', 'Ⅱ-1 수익 구조', 'Ⅰ-3 마진 구조',
                                  'Ⅰ-1 경계 (규제가 산업을 정의하는 경우)', 'Ⅰ-4 진입장벽 숫자')):
        cell(r.cells[1], v, 9)
    gap(doc, 3)

    para(doc, '인원별 배분 (권장)', 12, True, NAVY, after=3)
    t = table(doc, ['인원', '배분'], (3, W - 3), rows=3, first=['4명', '5명', '6명'])
    for r, v in zip(t.rows[1:], (
            'A: Ⅰ-1 · Ⅰ-2    B: Ⅰ-3 · Ⅱ-1    C: Ⅰ-4 · Ⅲ    D: Ⅰ-5 · Ⅱ-2    부록 · Ⅳ는 전원',
            'A: Ⅰ-1 · Ⅰ-2    B: Ⅰ-3 · Ⅱ-1    C: Ⅰ-4    D: Ⅰ-5 · Ⅱ-2    E: Ⅲ · 부록    Ⅳ는 전원',
            'A: Ⅰ-1    B: Ⅰ-2    C: Ⅰ-3 · Ⅱ-1    D: Ⅰ-4    E: Ⅰ-5 · Ⅱ-2    F: Ⅲ · 부록    Ⅳ는 전원')):
        cell(r.cells[1], v, 9)
    gap(doc, 3)
    note(doc, 'Ⅰ-5 지난 사이클과 Ⅱ-2 사이클 변수는 같은 사람이 맡는다. 지난 국면을 가른 변수가 곧 사이클 변수 후보다.')
    gap(doc, 3)
    para(doc, '금지 사항', 12, True, RED, after=3)
    for ln in (
        '증권사 리포트를 논지의 근거로 사용하지 않는다. 목표주가를 인용하지 않는다.',
        '검증 불가능한 표현을 쓰지 않는다. 유망하다 · 성장성이 크다 · 저평가되어 있다.',
        '전망 표현을 쓰지 않는다. 전망한다 · 수혜가 예상된다 · 성장이 기대된다.',
        '그래프 3개 이상은 팀이 직접 그린 것이어야 한다. 캡처 이미지는 출처를 달되 3개에 포함하지 않는다.',
        '모든 수치에 출처와 기준일을 기재한다.',
    ):
        para(doc, f'·  {ln}', 10, after=2)


    # ── Ⅰ-1 정의와 경계 ─────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅰ-1', '산업 정의와 경계', '1분',
           '정의표 + 사업보고서 3개사 비교표',
           '사업보고서 II장 「산업의 특성」 「산업의 성장성」 · KRX 업종분류',
           ['이 산업은 무엇을 팔고, 누가 사는가.  고객은 그것을 왜 사고, 대체재는 무엇인가.',
            '이 산업의 경계는 어디인가.  무엇을 포함하고 무엇을 뺐는가, 그 이유는 무엇인가.',
            '같은 산업을 세 회사가 어떻게 다르게 설명하는가.'])
    t = table(doc, ['항목', '기재', '출처 · 기준일'], (5, 15, W - 20), rows=7,
              first=['무엇을 파는가', '누가 사는가 (최종 고객)', '시장 규모 — 글로벌', '시장 규모 — 국내',
                     '과거 5년 성장률 (실적치)', '포함하는 것', '제외하는 것과 이유'])
    gap(doc, 5)
    para(doc, '사업보고서 비교 — 3개사가 자기 산업을 어떻게 설명하는가', 10, True, after=3)
    table(doc, ['기업', '산업의 특성 (한 줄)', '성장 동인 (한 줄)', '세 회사가 다르게 말하는 부분'],
          (4, 7.5, 7.5, W - 19), rows=3)

    # ── Ⅰ-2 밸류체인 ────────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅰ-2', '밸류체인', '1.5분',
           '밸류체인 그림 1개 + 단계별 기업 배치표',
           '사업보고서 II장 「주요 제품」 「원재료」 「매출처」 — 한 회사의 매출처가 다음 단계다',
           ['원재료에서 최종 소비자까지 몇 단계를 거치는가.  각 단계에서 무엇이 더해지는가.',
            '우리 담당 종목은 어느 단계에 몰려 있고, 어느 단계가 비어 있는가.',
            '비어 있는 단계는 누가 채우고 있는가 (해외 기업 · 비상장).'])
    box(doc, '밸류체인 그림을 여기에 붙인다  —  단계를 화살표로 잇고, 각 단계 아래에 대표 기업을 적는다', 5.2)
    gap(doc, 4)
    table(doc, ['단계', '하는 일', '국내 대표 기업', '해외 대표 기업', '우리 담당 종목', 'BM 비중 합'],
          (3.2, 6, 5, 5, 5, W - 24.2), rows=5)
    gap(doc, 3)
    note(doc, 'KRX 300에 편입 종목이 없는 단계도 그린다. 비어 있는 단계가 한국 시장에서 이 산업의 구조를 보여준다.')

    # ── Ⅰ-3 마진 구조 ───────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅰ-3', '마진은 어디에 쌓이는가', '1분',
           '단계별 영업이익률 5년 표 + 막대그래프 1개 + 교섭력 표',
           'DART 재무제표 (5개년) · 사업보고서 II장 「매출처」 「주요 원재료」',
           ['밸류체인에서 영업이익률이 가장 높은 단계는 어디이고, 가장 낮은 단계는 어디인가.',
            '고객에게 가격을 올릴 수 있는 단계와 원가 상승을 떠안는 단계는 각각 어디인가.',
            '영업이익률이 가장 크게 흔들리는 단계는 어디인가.'])
    table(doc, ['단계', '대표 기업', '영업이익률 5년 평균', '최저 (연도)', '최고 (연도)', '출처'],
          (3.2, 5, 4.5, 4, 4, W - 20.7), rows=4)
    gap(doc, 4)
    t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
    cell(t.rows[0].cells[0], '단계별 영업이익률 막대그래프  (5년 평균 · 최저~최고 범위 표시)', 9, False, GRAY,
         WD_ALIGN_PARAGRAPH.CENTER)
    c = t.rows[0].cells[1]; c.text = ''
    for i, ln in enumerate(('교섭력 판단 기준',
                            '·  상위 5대 고객 매출 비중이 50%를 넘으면 고객이 가격을 쥔다',
                            '·  원재료비가 매출원가의 60%를 넘으면 원가를 떠안는다',
                            '·  둘 다 해당하면 양쪽에서 눌리는 단계다',
                            '기준값은 산업마다 다르다. 단계끼리 비교해서 판단한다.')):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        style(p.add_run(ln), 9, i == 0, NAVY if i == 0 else (GRAY if i == 4 else None))
    widths(t, (15, W - 15)); t.rows[0].height = Cm(3.4)
    gap(doc, 4)
    table(doc, ['기업', '상위 1대 고객 비중', '상위 5대 고객 비중', '원재료비 / 매출원가', '주요 원재료', '출처'],
          (4, 4, 4, 4.5, 4.5, W - 21), rows=3)

    # ── Ⅰ-4 경쟁 구도 ───────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅰ-4', '경쟁 구도', '1분',
           '점유율 표 (5년 전 대비) + 진입장벽 숫자 하나',
           '사업보고서 II장 「시장점유율」 「경쟁요소」 「회사의 경쟁우위요소」 · 협회 통계',
           ['Ⅰ-3에서 마진이 가장 높았던 단계에서 누가 이기고 있는가.',
            '5년 사이 점유율이 누구에게서 누구에게로 옮겨 갔는가.',
            '새로 들어오려는 회사는 무엇에서 막히는가.  그것을 숫자로 보일 수 있는가.'])
    para(doc, '기준 시장 정의 (글로벌 / 국내, 금액 / 수량) : ______________________________________________',
         9, color=GRAY, after=4)
    table(doc, ['순위', '기업', '점유율 현재', '5년 전', '변화', '국가 · 상장 여부', '출처 · 기준일'],
          (1.8, 5, 3, 3, 2.5, 4.5, W - 19.8), rows=5)
    gap(doc, 5)
    table(doc, ['진입장벽 (한 문장)', '뒷받침하는 숫자 하나', '출처'], (11, 10, W - 21), rows=1)
    gap(doc, 3)
    note(doc, '설비 투자액 · 인증 소요기간 · 고객 전환비용 · 특허 건수 중 하나를 고른다. '
              '숫자가 없으면 진입장벽이 있다고 쓰지 않는다.')

    # ── Ⅰ-5 지난 사이클 ─────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅰ-5', '지난 사이클', '0.5분',
           '과거 10~15년 국면 타임라인 표',
           '사업보고서 II장 「경기변동의 특성」 · 대표사 과거 영업이익 (DART) · 협회 통계',
           ['이 산업은 지난 10~15년간 몇 번의 호황과 불황을 겪었는가.',
            '각 국면을 시작하고 끝낸 것은 무엇이었는가.',
            '같은 변수가 여러 국면에서 반복해서 등장하는가.'])
    table(doc, ['국면', '기간', '무엇이 일어났나', '계기가 된 변수', '지표 수치 (시작 → 끝)', '대표사 영업이익 변화', '출처'],
          (2.4, 3, 6, 4, 4.3, 4, W - 23.7), rows=5)
    gap(doc, 4)
    note(doc, '「지금이 어느 국면과 닮았는가」는 쓰지 않는다. 여기서 반복해서 등장한 변수가 Ⅱ-2 사이클 변수의 후보다.')

    # ── Ⅱ-1 수익 구조 ───────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅱ-1', '돈이 어디서 나오는가 — 수익 구조', '1.5분',
           '대표 3사 제품별 · 지역별 매출 구성표',
           '사업보고서 II장 「주요 제품 등의 현황」 「매출실적」',
           ['대표 3사의 매출은 어떤 제품과 어느 지역에서 나오는가.',
            '이 산업의 매출을 움직이는 것은 단가인가, 물량인가.'])
    para(doc, '제품별 매출 구성  (최근 사업연도 · 단위 %)', 10, True, after=3)
    table(doc, ['기업', '제품 1', '비중', '제품 2', '비중', '제품 3', '비중', '기타', '출처 · 기준일'],
          (3.6, 3.4, 1.8, 3.4, 1.8, 3.4, 1.8, 1.8, W - 21), rows=3)
    gap(doc, 4)
    para(doc, '지역별 매출 구성  (단위 %)', 10, True, after=3)
    table(doc, ['기업', '국내', '중국', '미국', '유럽', '기타', '수출 비중', '출처 · 기준일'],
          (3.6, 2.6, 2.6, 2.6, 2.6, 2.6, 3, W - 19.6), rows=3)
    gap(doc, 4)
    table(doc, ['단가 / 물량 / 둘 다', '근거 숫자'], (7, W - 7), rows=1)

    # ── Ⅱ-2 사이클 변수 ─────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅱ-2', '사이클을 움직이는 변수 3개', '2분',
           '5년 시계열 그래프 3개 직접 작성',
           'Ⅰ-5에서 반복 등장한 변수 · ECOS · 무역협회 K-stat · 업종 협회 통계',
           ['Ⅰ-5에서 반복해서 등장한 변수 중 무엇을 골랐고, 왜 골랐는가.',
            '세 변수와 대표사 영업이익률을 겹쳐 그리면 어떤 순서로 움직이는가.'])
    table(doc, ['#', '변수', '고른 이유', '데이터 소스', '기간', '최근값 · 기준일'],
          (1.2, 5, 8, 5.5, 3, W - 22.7), rows=3)
    gap(doc, 4)
    t = doc.add_table(rows=1, cols=3); t.style = 'Table Grid'
    for i, lab in enumerate(('변수 1 + 대표사 영업이익률', '변수 2 + 대표사 영업이익률', '변수 3 + 대표사 영업이익률')):
        cell(t.rows[0].cells[i], lab, 9, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)
    widths(t, (W / 3,) * 3); t.rows[0].height = Cm(5.3)
    gap(doc, 4)
    table(doc, ['관찰 (두 줄 · 전망이 아니라 관찰)'], (W,), rows=1)

    # ── Ⅲ 시장 기대 ─────────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅲ', '시장이 지금 기대하고 있는 것', '1.5분',
           '산포도 1개 + 컨센서스 변화율 표',
           '팀 스크리닝 시트 · FnGuide 컨센서스',
           ['시장은 이 섹터의 어느 종목에 가장 높은 배수를 주고 있는가.',
            'Ⅰ-3에서 마진이 높았던 단계의 종목이 배수도 높은가.'])
    box(doc, '산포도  —  가로축 12M Fwd PER, 세로축 ROE, 종목명 라벨 포함. 밸류체인 단계별로 색을 나눈다', 5.8)
    gap(doc, 4)
    table(doc, ['구분', '종목', '밸류체인 단계', '1년 전 전망', '현재 전망', '변화율', '12M Fwd PER', '출처 · 기준일'],
          (2.6, 4, 3.6, 3.2, 3.2, 2.6, 3, W - 22.2), rows=4,
          first=['상향 1', '상향 2', '하향 1', '하향 2'])

    # ── Ⅳ 모르는 것 ─────────────────────────────────────────────
    page(doc)
    header(doc, 'Ⅳ', '우리가 모르는 것 5개', '2분',
           '미결 질문 5개 — 10월 6일 종목 발제의 출발점',
           '이 문서 전체 · 팀 토의',
           ['이번 공부로 답이 나오지 않은 질문은 무엇인가.',
            '그것은 자료가 없어서 모르는 것인가, 자료는 있는데 아직 못 본 것인가.'])
    t = table(doc, ['#', '모르는 것 (질문 형태로)', '관련 장', '무엇을 보면 답이 나오는가', '자료 없음 / 미확인', '확인 담당'],
              (1, 8.5, 2.2, 8, 3.5, W - 23.2), rows=5)
    for i in range(1, 6):
        cell(t.rows[i].cells[0], str(i), 9, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 6)
    para(doc, '5일 공부로 산업을 다 알 수는 없다. 무엇을 모르는지 정확히 아는 것이 이 과제의 목표다.',
         10, True, NAVY, after=3)
    note(doc, '「자료 없음」은 그 섹터의 구조적 한계이고, 「미확인」은 다음 주 숙제다.')

    # ── 부록 용어집 ─────────────────────────────────────────────
    page(doc)
    para(doc, '부록   용어집 20개', 16, True, NAVY, after=1)
    para(doc, '발표 제외  |  섹터 전문용어 20개를 팀이 직접 정의한다. 신입이 읽고 이해할 수 있어야 한다',
         9, color=GRAY, after=6)
    t = doc.add_table(rows=11, cols=6); t.style = 'Table Grid'
    for col in (0, 3):
        cell(t.rows[0].cells[col], '#', 9, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, 'D6DCE4')
        cell(t.rows[0].cells[col + 1], '용어', 9, True, NAVY, None, 'D6DCE4')
        cell(t.rows[0].cells[col + 2], '정의 (한 줄)', 9, True, NAVY, None, 'D6DCE4')
    for i in range(1, 11):
        cell(t.rows[i].cells[0], str(i), 9, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell(t.rows[i].cells[3], str(i + 10), 9, align=WD_ALIGN_PARAGRAPH.CENTER)
    widths(t, (1, 3.5, 8.85, 1, 3.5, 8.85))



def main():
    doc = Document()
    build(doc)
    doc.save(OUT)
    print(f'저장: {OUT}')


if __name__ == '__main__':
    main()
